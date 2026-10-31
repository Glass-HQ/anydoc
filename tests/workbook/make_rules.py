import sys
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule, Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
wb = Workbook(); ws = wb.active; ws.title = "Rules"
ws.append(["Name", "Score", "Status", "Heat", "Bar"])
rows = [("Ada", 92, "Done", 10, 5), ("Alan", 45, "Late", 50, 20), ("Grace", 78, "Done", 90, 80), ("Linus", 30, "Late", 30, 100), ("Ada", 60, "Open", 70, 50)]
for r in rows: ws.append(r)
red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
ws.conditional_formatting.add("B2:B6", CellIsRule(operator="greaterThan", formula=["70"], fill=PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"), font=Font(bold=True, color="006100")))
ws.conditional_formatting.add("C2:C6", Rule(type="containsText", operator="containsText", text="Late", dxf=DifferentialStyle(fill=red)))
ws.conditional_formatting.add("D2:D6", ColorScaleRule(start_type="min", start_color="F8696B", end_type="max", end_color="63BE7B"))
ws.conditional_formatting.add("E2:E6", DataBarRule(start_type="min", end_type="max", color="638EC6"))
ws.conditional_formatting.add("A2:A6", Rule(type="duplicateValues", dxf=DifferentialStyle(font=Font(italic=True, color="9C0006"))))
t = Table(displayName="People", ref="A1:E6"); t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True); ws.add_table(t)
dv = DataValidation(type="list", formula1='"Done,Late,Open"', allow_blank=True); ws.add_data_validation(dv); dv.add("C2:C6")
ws2 = wb.create_sheet("Lists"); ws2.append(["Red"]); ws2.append(["Green"]); ws2.append(["Blue"])
dv2 = DataValidation(type="list", formula1="Lists!$A$1:$A$3"); ws.add_data_validation(dv2); dv2.add("F2")
wb.save(sys.argv[1])
